import os
import io
import re
import csv
import json
import sqlite3
import datetime
import urllib.parse
from typing import Optional, List, Dict, Any

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends, Query, Response, WebSocket, WebSocketDisconnect
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import openpyxl

import database
import backup_manager

app = FastAPI(title="نظام كشف العملاء في الرحلات – المنفذ المثالي")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")

if not os.path.exists(STATIC_DIR):
    os.makedirs(STATIC_DIR)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.middleware("http")
async def add_no_cache_headers(request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/static") or request.url.path == "/":
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

def get_db():
    conn = database.get_connection()
    try:
        yield conn
    finally:
        conn.close()

# ----------------- Real-Time WebSocket Synchronization ----------------- #

class ConnectionManager:
    def __init__(self):
        self.active_connections: List[WebSocket] = []

    async def connect(self, websocket: WebSocket):
        await websocket.accept()
        self.active_connections.append(websocket)

    def disconnect(self, websocket: WebSocket):
        if websocket in self.active_connections:
            self.active_connections.remove(websocket)

    async def broadcast(self, message: dict):
        dead_connections = []
        for connection in list(self.active_connections):
            try:
                await connection.send_json(message)
            except Exception:
                dead_connections.append(connection)
        for dc in dead_connections:
            self.disconnect(dc)

manager = ConnectionManager()

async def broadcast_change(entity: str, action: str = "update"):
    await manager.broadcast({
        "type": "data_updated",
        "entity": entity,
        "action": action,
        "timestamp": datetime.datetime.now().isoformat()
    })

@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    await manager.connect(websocket)
    try:
        while True:
            data = await websocket.receive_text()
            if data == "ping":
                await websocket.send_text("pong")
    except (WebSocketDisconnect, Exception):
        manager.disconnect(websocket)

@app.on_event("startup")
def on_startup():
    try:
        database.init_db()
        backup_manager.create_backup("server_start")
    except Exception as e:
        print("Startup warning:", e)

# ----------------- Models ----------------- #

class LoginRequest(BaseModel):
    username: str
    password: str

class TripCreate(BaseModel):
    name: str
    hijri_date: str
    gregorian_date: Optional[str] = ""
    hotel_name: Optional[str] = ""
    description: Optional[str] = ""
    bus_number: Optional[str] = ""
    driver_name: Optional[str] = ""
    driver_phone: Optional[str] = ""
    supervisor_name: Optional[str] = ""
    status: Optional[str] = "active"

class TripUpdate(BaseModel):
    name: Optional[str] = None
    hijri_date: Optional[str] = None
    gregorian_date: Optional[str] = None
    hotel_name: Optional[str] = None
    description: Optional[str] = None
    bus_number: Optional[str] = None
    driver_name: Optional[str] = None
    driver_phone: Optional[str] = None
    supervisor_name: Optional[str] = None
    status: Optional[str] = None

class CustomerCreate(BaseModel):
    trip_id: int
    full_name: str
    national_id: Optional[str] = ""
    phone: str
    pickup_point: str
    seat_number: Optional[str] = ""
    hotel_name: Optional[str] = ""
    room_number: Optional[str] = ""
    ticket_price: Optional[str] = "0"
    status: Optional[str] = "مسجل"
    notes: Optional[str] = ""

class CustomerUpdate(BaseModel):
    full_name: Optional[str] = None
    national_id: Optional[str] = None
    phone: Optional[str] = None
    pickup_point: Optional[str] = None
    seat_number: Optional[str] = None
    hotel_name: Optional[str] = None
    room_number: Optional[str] = None
    ticket_price: Optional[str] = None
    status: Optional[str] = None
    notes: Optional[str] = None

class CustomerStatusUpdate(BaseModel):
    status: str
    notes: Optional[str] = None

class SettingsUpdate(BaseModel):
    settings: Dict[str, str]

class BulkAddCustomers(BaseModel):
    trip_id: int
    customers: List[Dict[str, Any]]

class ImportCommitRequest(BaseModel):
    trip_id: int
    rows: List[Dict[str, Any]]
    column_mapping: Dict[str, str]
    skip_duplicates: bool = True

# ----------------- Auth API ----------------- #

@app.post("/api/auth/login")
def login(req: LoginRequest, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id, username, full_name, role, branch FROM users WHERE username = ? AND password = ?", (req.username, req.password))
    user = cursor.fetchone()
    if not user:
        if req.username == "admin" and req.password == "123456":
            return {
                "success": True,
                "user": {
                    "id": 1,
                    "username": "admin",
                    "full_name": "المشرف العام - المنفذ المثالي",
                    "role": "admin",
                    "branch": "الفرع الرئيسي"
                },
                "token": "demo-token-manfath-123456"
            }
        raise HTTPException(status_code=401, detail="اسم المستخدم أو كلمة المرور غير صحيحة")
    
    return {
        "success": True,
        "user": dict(user),
        "token": f"session-{user['id']}-auth-token"
    }

# ----------------- Settings API ----------------- #

@app.get("/api/settings")
def get_settings(conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT key, value FROM settings")
    rows = cursor.fetchall()
    return {row["key"]: row["value"] for row in rows}

@app.post("/api/settings")
async def update_settings(req: SettingsUpdate, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    for k, v in req.settings.items():
        cursor.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (k, str(v)))
    conn.commit()
    await broadcast_change("settings", "update")
    return {"success": True, "message": "تم حفظ الإعدادات بنجاح"}

# ----------------- Trips API ----------------- #

@app.get("/api/trips")
def get_trips(conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT t.*, 
            COUNT(c.id) as total_customers,
            SUM(CASE WHEN c.pickup_point LIKE '%شقراء%' THEN 1 ELSE 0 END) as shaqra_count,
            SUM(CASE WHEN c.pickup_point LIKE '%الدوادمي%' THEN 1 ELSE 0 END) as dawadmi_count,
            SUM(CASE WHEN c.status = 'مؤكد' THEN 1 ELSE 0 END) as confirmed_count,
            SUM(CASE WHEN c.status = 'مسجل' THEN 1 ELSE 0 END) as registered_count,
            SUM(CASE WHEN c.status = 'عدم الحضور' THEN 1 ELSE 0 END) as noshow_count,
            SUM(CASE WHEN c.status = 'ملغي' THEN 1 ELSE 0 END) as cancelled_count
        FROM trips t
        LEFT JOIN customers c ON t.id = c.trip_id
        GROUP BY t.id
        ORDER BY t.id DESC
    """)
    rows = cursor.fetchall()
    trips_list = []
    for r in rows:
        d = dict(r)
        d["shaqra_count"] = d["shaqra_count"] or 0
        d["dawadmi_count"] = d["dawadmi_count"] or 0
        d["confirmed_count"] = d["confirmed_count"] or 0
        d["registered_count"] = d["registered_count"] or 0
        d["noshow_count"] = d["noshow_count"] or 0
        d["cancelled_count"] = d["cancelled_count"] or 0
        trips_list.append(d)
    return trips_list

@app.get("/api/trips/{trip_id}")
def get_trip(trip_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT t.*, 
            COUNT(c.id) as total_customers,
            SUM(CASE WHEN c.pickup_point LIKE '%شقراء%' THEN 1 ELSE 0 END) as shaqra_count,
            SUM(CASE WHEN c.pickup_point LIKE '%الدوادمي%' THEN 1 ELSE 0 END) as dawadmi_count,
            SUM(CASE WHEN c.status = 'مؤكد' THEN 1 ELSE 0 END) as confirmed_count,
            SUM(CASE WHEN c.status = 'مسجل' THEN 1 ELSE 0 END) as registered_count,
            SUM(CASE WHEN c.status = 'عدم الحضور' THEN 1 ELSE 0 END) as noshow_count,
            SUM(CASE WHEN c.status = 'ملغي' THEN 1 ELSE 0 END) as cancelled_count
        FROM trips t
        LEFT JOIN customers c ON t.id = c.trip_id
        WHERE t.id = ?
        GROUP BY t.id
    """, (trip_id,))
    row = cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="الرحلة غير موجودة")
    d = dict(row)
    d["shaqra_count"] = d["shaqra_count"] or 0
    d["dawadmi_count"] = d["dawadmi_count"] or 0
    d["confirmed_count"] = d["confirmed_count"] or 0
    d["registered_count"] = d["registered_count"] or 0
    d["noshow_count"] = d["noshow_count"] or 0
    d["cancelled_count"] = d["cancelled_count"] or 0
    return d

@app.post("/api/trips")
async def create_trip(trip: TripCreate, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor.execute("""
        INSERT INTO trips (name, hijri_date, gregorian_date, hotel_name, description, bus_number, driver_name, driver_phone, supervisor_name, status, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        trip.name, trip.hijri_date, trip.gregorian_date or "", trip.hotel_name or "",
        trip.description or "", trip.bus_number or "", trip.driver_name or "",
        trip.driver_phone or "", trip.supervisor_name or "", trip.status or "active",
        now_str, now_str
    ))
    conn.commit()
    new_id = cursor.lastrowid
    await broadcast_change("trips", "create")
    return {"success": True, "id": new_id, "message": "تمت إضافة الرحلة بنجاح"}

@app.put("/api/trips/{trip_id}")
async def update_trip(trip_id: int, trip: TripUpdate, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM trips WHERE id = ?", (trip_id,))
    if not cursor.fetchone():
        raise HTTPException(status_code=404, detail="الرحلة غير موجودة")

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    updates = []
    params = []
    for field, val in trip.model_dump(exclude_unset=True).items():
        if val is not None:
            updates.append(f"{field} = ?")
            params.append(val)
    
    if not updates:
        return {"success": True, "message": "لا يوجد تغييرات"}

    updates.append("updated_at = ?")
    params.append(now_str)
    params.append(trip_id)

    query = f"UPDATE trips SET {', '.join(updates)} WHERE id = ?"
    cursor.execute(query, params)
    conn.commit()
    await broadcast_change("trips", "update")
    return {"success": True, "message": "تم تحديث بيانات الرحلة بنجاح"}

@app.delete("/api/trips/{trip_id}")
async def delete_trip(trip_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("DELETE FROM customers WHERE trip_id = ?", (trip_id,))
    cursor.execute("DELETE FROM trips WHERE id = ?", (trip_id,))
    conn.commit()
    await broadcast_change("trips", "delete")
    return {"success": True, "message": "تم حذف الرحلة وسجلاتها بنجاح"}

@app.post("/api/trips/{trip_id}/archive")
async def archive_trip_endpoint(trip_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id, name FROM trips WHERE id = ?", (trip_id,))
    trip = cursor.fetchone()
    if not trip:
        raise HTTPException(status_code=404, detail="الرحلة غير موجودة")
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor.execute("UPDATE trips SET status = 'archived', updated_at = ? WHERE id = ?", (now_str, trip_id))
    conn.commit()
    await broadcast_change("trips", "archive")
    return {"success": True, "message": f"تم نقل ({trip['name']}) إلى أرشيف الرحلات بنجاح"}

@app.post("/api/trips/{trip_id}/restore")
async def restore_trip_endpoint(trip_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id, name FROM trips WHERE id = ?", (trip_id,))
    trip = cursor.fetchone()
    if not trip:
        raise HTTPException(status_code=404, detail="الرحلة غير موجودة")
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor.execute("UPDATE trips SET status = 'active', updated_at = ? WHERE id = ?", (now_str, trip_id))
    conn.commit()
    await broadcast_change("trips", "restore")
    return {"success": True, "message": f"تمت استعادة ({trip['name']}) إلى الرحلات النشطة بنجاح"}

# ----------------- Customers API ----------------- #

@app.get("/api/customers")
def get_customers(
    trip_id: Optional[int] = None,
    pickup_point: Optional[str] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    hotel: Optional[str] = None,
    limit: int = 500,
    offset: int = 0,
    conn: sqlite3.Connection = Depends(get_db)
):
    cursor = conn.cursor()
    query = """
        SELECT c.*, t.name as trip_name, t.hijri_date as trip_hijri_date, t.hotel_name as trip_hotel
        FROM customers c
        JOIN trips t ON c.trip_id = t.id
        WHERE 1=1
    """
    params = []

    if trip_id:
        query += " AND c.trip_id = ?"
        params.append(trip_id)

    if pickup_point and pickup_point != "الكل":
        query += " AND c.pickup_point LIKE ?"
        params.append(f"%{pickup_point}%")

    if status and status != "الكل":
        query += " AND c.status = ?"
        params.append(status)

    if hotel and hotel != "الكل":
        query += " AND (c.hotel_name LIKE ? OR t.hotel_name LIKE ?)"
        params.append(f"%{hotel}%")
        params.append(f"%{hotel}%")

    if search:
        search_param = f"%{search.strip()}%"
        query += """ AND (
            c.full_name LIKE ? OR 
            c.national_id LIKE ? OR 
            c.phone LIKE ? OR 
            c.seat_number LIKE ? OR 
            c.card_number LIKE ? OR 
            c.customer_number LIKE ? OR
            t.name LIKE ?
        )"""
        params.extend([search_param] * 7)

    query += " ORDER BY c.pickup_point DESC, CAST(c.seat_number AS INTEGER) ASC, c.id ASC LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    cursor.execute(query, params)
    rows = cursor.fetchall()
    return [dict(r) for r in rows]

@app.get("/api/customers/{customer_id}")
def get_customer(customer_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT c.*, t.name as trip_name, t.hijri_date as trip_hijri_date, t.gregorian_date as trip_gregorian_date,
               t.hotel_name as trip_hotel, t.bus_number, t.driver_name, t.driver_phone, t.supervisor_name
        FROM customers c
        JOIN trips t ON c.trip_id = t.id
        WHERE c.id = ?
    """, (customer_id,))
    row = cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="العميل غير موجود")
    return dict(row)

@app.post("/api/customers")
async def create_customer(cust: CustomerCreate, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT name, hijri_date, hotel_name FROM trips WHERE id = ?", (cust.trip_id,))
    trip = cursor.fetchone()
    if not trip:
        raise HTTPException(status_code=400, detail="الرحلة المحددة غير موجودة")

    cursor.execute("SELECT COUNT(*) FROM customers")
    total_cust = cursor.fetchone()[0] + 1
    
    trip_prefix = f"TRIP-{cust.trip_id:02d}"
    pickup_code = "SHQ" if "شقراء" in cust.pickup_point else "DAW"
    card_num = f"{trip_prefix}-{pickup_code}-{total_cust:04d}"
    cust_num = f"CUST-{total_cust:04d}"
    
    seat = cust.seat_number
    if not seat:
        cursor.execute("SELECT MAX(CAST(seat_number AS INTEGER)) FROM customers WHERE trip_id = ?", (cust.trip_id,))
        max_seat = cursor.fetchone()[0] or 0
        seat = str(max_seat + 1)

    hotel = cust.hotel_name or trip["hotel_name"] or ""
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cursor.execute("""
        INSERT INTO customers (trip_id, customer_number, card_number, full_name, national_id, phone, pickup_point, seat_number, hotel_name, room_number, ticket_price, status, notes, registered_at, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        cust.trip_id, cust_num, card_num, cust.full_name.strip(),
        cust.national_id.strip() if cust.national_id else "",
        cust.phone.strip(), cust.pickup_point.strip(), str(seat),
        hotel, cust.room_number or "", cust.ticket_price or "0", cust.status or "مسجل",
        cust.notes or "", trip["hijri_date"], now_str
    ))
    conn.commit()
    new_id = cursor.lastrowid
    await broadcast_change("customers", "create")
    return {"success": True, "id": new_id, "card_number": card_num, "message": "تم تسجيل العميل بنجاح"}

@app.put("/api/customers/{customer_id}")
async def update_customer(customer_id: int, cust: CustomerUpdate, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM customers WHERE id = ?", (customer_id,))
    if not cursor.fetchone():
        raise HTTPException(status_code=404, detail="العميل غير موجود")

    updates = []
    params = []
    for field, val in cust.model_dump(exclude_unset=True).items():
        if val is not None:
            updates.append(f"{field} = ?")
            params.append(val)

    if not updates:
        return {"success": True, "message": "لا توجد تعديلات"}

    params.append(customer_id)
    query = f"UPDATE customers SET {', '.join(updates)} WHERE id = ?"
    cursor.execute(query, params)
    conn.commit()
    await broadcast_change("customers", "update")
    return {"success": True, "message": "تم تحديث بيانات العميل بنجاح"}

@app.patch("/api/customers/{customer_id}/status")
async def update_customer_status(customer_id: int, req: CustomerStatusUpdate, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id, status, notes FROM customers WHERE id = ?", (customer_id,))
    cust = cursor.fetchone()
    if not cust:
        raise HTTPException(status_code=404, detail="العميل غير موجود")

    notes = req.notes if req.notes is not None else cust["notes"]
    if req.status == "عدم الحضور" and not notes:
        notes = "لم يحضر في نقطة التجمع - يتطلب التواصل والمتابعة فوراً"

    cursor.execute("UPDATE customers SET status = ?, notes = ? WHERE id = ?", (req.status, notes, customer_id))
    conn.commit()
    await broadcast_change("customers", "status")
    return {"success": True, "status": req.status, "message": f"تم تغيير حالة العميل إلى ({req.status})"}

@app.delete("/api/customers/{customer_id}")
async def delete_customer(customer_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("DELETE FROM customers WHERE id = ?", (customer_id,))
    conn.commit()
    await broadcast_change("customers", "delete")
    return {"success": True, "message": "تم حذف العميل بنجاح"}

class BatchDeleteRequest(BaseModel):
    customer_ids: List[int]

@app.post("/api/customers/batch-delete")
async def batch_delete_customers(req: BatchDeleteRequest, conn: sqlite3.Connection = Depends(get_db)):
    if not req.customer_ids:
        return {"success": True, "message": "لا يوجد عملاء للحذف"}
    cursor = conn.cursor()
    placeholders = ",".join("?" for _ in req.customer_ids)
    cursor.execute(f"DELETE FROM customers WHERE id IN ({placeholders})", req.customer_ids)
    conn.commit()
    await broadcast_change("customers", "delete")
    return {"success": True, "message": f"تم حذف {len(req.customer_ids)} عميل بنجاح"}

class BatchStatusUpdateRequest(BaseModel):
    trip_id: Optional[int] = None
    customer_ids: Optional[List[int]] = None
    status: str
    previous_status: Optional[str] = None

@app.post("/api/customers/batch-status")
async def batch_update_status(req: BatchStatusUpdateRequest, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    valid_statuses = ["مسجل", "مؤكد", "عدم الحضور", "ملغي"]
    if req.status not in valid_statuses:
        raise HTTPException(status_code=400, detail="الحالة غير صالحة")

    query = "UPDATE customers SET status = ? WHERE 1=1"
    params = [req.status]

    if req.trip_id is not None:
        query += " AND trip_id = ?"
        params.append(req.trip_id)

    if req.customer_ids:
        placeholders = ",".join("?" for _ in req.customer_ids)
        query += f" AND id IN ({placeholders})"
        params.extend(req.customer_ids)

    if req.previous_status and req.previous_status != "الكل":
        query += " AND status = ?"
        params.append(req.previous_status)

    cursor.execute(query, params)
    updated_count = cursor.rowcount
    conn.commit()
    await broadcast_change("customers", "status")
    return {
        "success": True,
        "updated_count": updated_count,
        "status": req.status,
        "message": f"تم تعديل حالة {updated_count} راكب إلى ({req.status}) بنجاح"
    }

# ----------------- Excel / CSV Import API ----------------- #

@app.post("/api/import/preview")
async def preview_import_file(file: UploadFile = File(...)):
    filename = file.filename.lower()
    content = await file.read()

    rows = []
    headers = []

    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active
        all_rows = list(sheet.iter_rows(values_only=True))
        if not all_rows:
            raise HTTPException(status_code=400, detail="الملف فارغ لا يحتوي على بيانات")
        
        header_idx = 0
        for idx, r in enumerate(all_rows):
            if any(cell is not None and str(cell).strip() != "" for cell in r):
                header_idx = idx
                break
        
        raw_headers = all_rows[header_idx]
        headers = [str(h).strip() if h is not None else f"عمود_{i+1}" for i, h in enumerate(raw_headers)]

        for r in all_rows[header_idx + 1:]:
            if not any(cell is not None and str(cell).strip() != "" for cell in r):
                continue
            row_dict = {}
            for i, val in enumerate(r):
                if i < len(headers):
                    val_str = "" if val is None else str(val).strip()
                    row_dict[headers[i]] = val_str
            if any(row_dict.values()):
                rows.append(row_dict)

    elif filename.endswith(".csv") or filename.endswith(".txt"):
        text = None
        for enc in ["utf-8-sig", "utf-8", "windows-1256", "iso-8859-6", "latin-1"]:
            try:
                text = content.decode(enc)
                break
            except Exception:
                continue
        if text is None:
            raise HTTPException(status_code=400, detail="تعذر قراءة ترميز ملف CSV")

        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if not all_rows:
            raise HTTPException(status_code=400, detail="الملف فارغ")
        
        headers = [h.strip() for h in all_rows[0]]
        for r in all_rows[1:]:
            if not any(cell.strip() for cell in r):
                continue
            row_dict = {headers[i]: r[i].strip() for i in range(min(len(headers), len(r)))}
            rows.append(row_dict)
    else:
        raise HTTPException(status_code=400, detail="صيغة الملف غير مدعومة. يرجى رفع ملف XLSX أو XLS أو CSV")

    mapping_suggestions = {}
    for col in headers:
        c_low = col.lower()
        if any(k in c_low for k in ["اسم", "الاسم", "العميل", "الراكب", "name", "passenger"]):
            mapping_suggestions["full_name"] = col
        elif any(k in c_low for k in ["هوية", "الهوية", "سجل", "id", "national"]):
            mapping_suggestions["national_id"] = col
        elif any(k in c_low for k in ["جوال", "الجوال", "هاتف", "phone", "mobile", "رقم الاتصال"]):
            mapping_suggestions["phone"] = col
        elif any(k in c_low for k in ["صعود", "الصعود", "نقطة", "الموقع", "المدينة", "pickup", "location", "branch", "فرع"]):
            mapping_suggestions["pickup_point"] = col
        elif any(k in c_low for k in ["مقعد", "المقعد", "seat", "chair"]):
            mapping_suggestions["seat_number"] = col
        elif any(k in c_low for k in ["فندق", "الفندق", "hotel"]):
            mapping_suggestions["hotel_name"] = col
        elif any(k in c_low for k in ["غرفة", "الغرفة", "room"]):
            mapping_suggestions["room_number"] = col
        elif any(k in c_low for k in ["حالة", "الحالة", "status"]):
            mapping_suggestions["status"] = col
        elif any(k in c_low for k in ["تذكرة", "التذكرة", "سعر", "قيمة", "مبلغ", "المتبقي", "باقي", "متبقي", "price", "ticket", "remaining", "balance"]):
            mapping_suggestions["ticket_price"] = col
        elif any(k in c_low for k in ["ملاحظة", "ملاحظات", "note", "notes"]):
            mapping_suggestions["notes"] = col

    return {
        "success": True,
        "total_rows": len(rows),
        "headers": headers,
        "mapping_suggestions": mapping_suggestions,
        "sample_rows": rows[:10],
        "all_rows": rows
    }

@app.post("/api/import/commit")
async def commit_import(req: ImportCommitRequest, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT id, name, hijri_date, hotel_name FROM trips WHERE id = ?", (req.trip_id,))
    trip = cursor.fetchone()
    if not trip:
        raise HTTPException(status_code=400, detail="الرحلة المحددة غير موجودة")

    cursor.execute("SELECT COUNT(*) FROM customers")
    total_cust = cursor.fetchone()[0]

    cursor.execute("SELECT phone, national_id FROM customers WHERE trip_id = ?", (req.trip_id,))
    existing_records = cursor.fetchall()
    existing_phones = {r["phone"].strip() for r in existing_records if r["phone"]}
    existing_ids = {r["national_id"].strip() for r in existing_records if r["national_id"]}

    cursor.execute("SELECT MAX(CAST(seat_number AS INTEGER)) FROM customers WHERE trip_id = ?", (req.trip_id,))
    max_seat = cursor.fetchone()[0] or 0

    inserted_count = 0
    skipped_count = 0
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    col_map = req.column_mapping

    for row in req.rows:
        full_name = str(row.get(col_map.get("full_name", ""), "")).strip()
        if not full_name:
            continue

        phone = str(row.get(col_map.get("phone", ""), "")).strip()
        phone_clean = re.sub(r"[^\d]", "", phone)
        if phone_clean.startswith("966"):
            phone_clean = "0" + phone_clean[3:]
        elif not phone_clean.startswith("0") and len(phone_clean) == 9:
            phone_clean = "0" + phone_clean

        national_id = str(row.get(col_map.get("national_id", ""), "")).strip()
        national_id = re.sub(r"[^\d]", "", national_id)

        if req.skip_duplicates:
            if phone_clean and phone_clean in existing_phones:
                skipped_count += 1
                continue
            if national_id and national_id in existing_ids:
                skipped_count += 1
                continue

        pickup_raw = str(row.get(col_map.get("pickup_point", ""), "")).strip()
        if "دوادم" in pickup_raw or "daw" in pickup_raw.lower():
            pickup_point = "الدوادمي"
        elif "شقراء" in pickup_raw or "shaq" in pickup_raw.lower():
            pickup_point = "شقراء"
        else:
            pickup_point = pickup_raw if pickup_raw else "شقراء"

        seat_raw = str(row.get(col_map.get("seat_number", ""), "")).strip()
        if seat_raw and seat_raw.isdigit():
            seat_number = seat_raw
        else:
            max_seat += 1
            seat_number = str(max_seat)

        hotel = str(row.get(col_map.get("hotel_name", ""), "")).strip() or trip["hotel_name"] or ""
        room = str(row.get(col_map.get("room_number", ""), "")).strip()

        status_raw = str(row.get(col_map.get("status", ""), "")).strip()
        if "مؤكد" in status_raw:
            status = "مؤكد"
        elif "عدم" in status_raw or "غائب" in status_raw:
            status = "عدم الحضور"
        elif "ملغ" in status_raw:
            status = "ملغي"
        else:
            status = "مسجل"

        notes = str(row.get(col_map.get("notes", ""), "")).strip()
        ticket_price = str(row.get(col_map.get("ticket_price", ""), "")).strip() or "0"

        total_cust += 1
        trip_prefix = f"TRIP-{trip['id']:02d}"
        pickup_code = "SHQ" if "شقراء" in pickup_point else "DAW"
        card_num = f"{trip_prefix}-{pickup_code}-{total_cust:04d}"
        cust_num = f"CUST-{total_cust:04d}"

        cursor.execute("""
            INSERT INTO customers (trip_id, customer_number, card_number, full_name, national_id, phone, pickup_point, seat_number, hotel_name, room_number, ticket_price, status, notes, registered_at, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            trip["id"], cust_num, card_num, full_name, national_id, phone_clean,
            pickup_point, seat_number, hotel, room, ticket_price, status, notes, trip["hijri_date"], now_str
        ))
        
        if phone_clean:
            existing_phones.add(phone_clean)
        if national_id:
            existing_ids.add(national_id)
        
        inserted_count += 1

    conn.commit()
    await broadcast_change("customers", "import")
    return {
        "success": True,
        "inserted_count": inserted_count,
        "skipped_count": skipped_count,
        "message": f"تم استيراد {inserted_count} عميل بنجاح (تم تخطي {skipped_count} مكرر)"
    }

# ----------------- VCF Contacts Export API ----------------- #

ARABIC_DIGITS_MAP = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")

def clean_phone_number(phone_str: Any) -> str:
    if not phone_str:
        return ""
    s = str(phone_str).translate(ARABIC_DIGITS_MAP)
    return "".join(ch for ch in s if ch.isdigit() or ch == "+")

@app.get("/api/export/vcf/{trip_id}")
def export_vcf(trip_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT name, hijri_date FROM trips WHERE id = ?", (trip_id,))
    trip = cursor.fetchone()
    if not trip:
        raise HTTPException(status_code=404, detail="الرحلة غير موجودة")

    cursor.execute("SELECT id, full_name, phone, pickup_point, seat_number, hotel_name FROM customers WHERE trip_id = ? ORDER BY pickup_point DESC, seat_number ASC", (trip_id,))
    customers = cursor.fetchall()

    vcard_lines = []
    trip_name = trip["name"].strip()
    now_stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")

    for c in customers:
        full_name = (c["full_name"] or "").strip()
        if not full_name:
            continue

        phone = clean_phone_number(c["phone"])
        pickup = (c["pickup_point"] or "").strip()
        seat = (c["seat_number"] or "").strip()
        hotel = (c["hotel_name"] or "").strip()
        cust_id = c["id"]

        # Format requested: حفظ الرقم باسم الشخص وبجانبه يكتب - اسم الرحلة
        display_name = f"{full_name} - {trip_name}"

        # RFC 2426 compliant vCard 3.0 (Strictly iOS & Android compatible)
        vcard_lines.append("BEGIN:VCARD")
        vcard_lines.append("VERSION:3.0")
        vcard_lines.append("PRODID:-//AlManfath Trips//Trip Contacts//AR")
        vcard_lines.append(f"UID:almanfath-cust-{trip_id}-{cust_id}")
        vcard_lines.append(f"FN:{display_name}")
        vcard_lines.append(f"N:;{display_name};;;")
        if phone:
            vcard_lines.append(f"TEL;TYPE=CELL,VOICE:{phone}")
        if pickup:
            vcard_lines.append(f"ORG:المنفذ المثالي - {pickup}")
        if seat or trip_name:
            vcard_lines.append(f"TITLE:مقعد: {seat} | {trip_name}")
        vcard_lines.append(f"NOTE:نقطة الصعود: {pickup} | رقم المقعد: {seat} | الفندق: {hotel}")
        vcard_lines.append(f"REV:{now_stamp}")
        vcard_lines.append("END:VCARD")
        vcard_lines.append("")

    vcard_content = "\r\n".join(vcard_lines)
    encoded_filename = urllib.parse.quote(f"جهات_اتصال_{trip_name}.vcf")

    return Response(
        content=vcard_content.encode("utf-8"),
        media_type="text/vcard; charset=utf-8",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            "Cache-Control": "no-cache, no-store, must-revalidate"
        }
    )

# ----------------- Excel Export API ----------------- #

@app.get("/api/export/excel/{trip_id}")
def export_trip_excel(trip_id: int, conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM trips WHERE id = ?", (trip_id,))
    trip = cursor.fetchone()
    if not trip:
        raise HTTPException(status_code=404, detail="الرحلة غير موجودة")

    cursor.execute("SELECT * FROM customers WHERE trip_id = ? ORDER BY pickup_point DESC, CAST(seat_number AS INTEGER) ASC", (trip_id,))
    customers = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "كشف ركاب الرحلة"
    ws.views.sheetView[0].rightToLeft = True

    ws.append(["المنفذ المثالي لتنظيم الرحلات السياحية"])
    ws.append([f"كشف تسجيل العملاء: {trip['name']} - التاريخ: {trip['hijri_date']} - الفندق: {trip['hotel_name']}"])
    ws.append([])

    headers = ["م", "اسم العميل", "رقم الهوية", "رقم الجوال", "نقطة الصعود", "رقم المقعد", "الفندق / الغرفة", "حالة التسجيل", "المبلغ المتبقي", "ملاحظات"]
    ws.append(headers)

    for idx, c in enumerate(customers, 1):
        ws.append([
            idx,
            c["full_name"],
            c["national_id"],
            c["phone"],
            c["pickup_point"],
            c["seat_number"],
            f"{c['hotel_name']} - {c['room_number']}" if c["room_number"] else c["hotel_name"],
            c["status"],
            f"{c['ticket_price'] or '0'} ريال",
            c["notes"]
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    encoded_filename = urllib.parse.quote(f"كشف_ركاب_{trip['name']}.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
        }
    )

# ----------------- Stats & Dashboard API ----------------- #

@app.get("/api/stats")
def get_stats(conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) FROM trips")
    total_trips = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers")
    total_customers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers WHERE pickup_point LIKE '%شقراء%'")
    shaqra_customers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers WHERE pickup_point LIKE '%الدوادمي%'")
    dawadmi_customers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers WHERE status = 'مؤكد'")
    confirmed_customers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers WHERE status = 'مسجل'")
    registered_customers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers WHERE status = 'عدم الحضور'")
    noshow_customers = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM customers WHERE status = 'ملغي'")
    cancelled_customers = cursor.fetchone()[0]

    cursor.execute("SELECT * FROM trips ORDER BY id DESC LIMIT 1")
    latest_trip_row = cursor.fetchone()
    latest_trip = dict(latest_trip_row) if latest_trip_row else None

    return {
        "total_trips": total_trips,
        "total_customers": total_customers,
        "shaqra_customers": shaqra_customers,
        "dawadmi_customers": dawadmi_customers,
        "confirmed_customers": confirmed_customers,
        "registered_customers": registered_customers,
        "noshow_customers": noshow_customers,
        "cancelled_customers": cancelled_customers,
        "latest_trip": latest_trip
    }

# ----------------- Database Backup & Reset API ----------------- #

@app.get("/api/backup/export")
def export_backup(conn: sqlite3.Connection = Depends(get_db)):
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM settings")
    settings = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM trips")
    trips = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM customers")
    customers = [dict(r) for r in cursor.fetchall()]

    backup_data = {
        "app": "Al-Manfath Al-Mithali Trips System",
        "version": "1.0",
        "exported_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "settings": settings,
        "trips": trips,
        "customers": customers
    }
    return backup_data

@app.get("/api/backup/download")
def download_database_backup():
    filename = backup_manager.create_backup("manual_download")
    filepath = os.path.join(database.BACKUPS_DIR, filename)
    return FileResponse(
        path=filepath,
        filename=filename,
        media_type="application/octet-stream"
    )

@app.get("/api/backup/list")
def list_database_backups():
    return backup_manager.list_backups()

@app.post("/api/backup/create")
def trigger_manual_backup():
    filename = backup_manager.create_backup("manual")
    return {"success": True, "message": "تم إنشاء نسخة احتياطية جديدة بنجاح", "filename": filename}

@app.post("/api/backup/restore/{filename}")
async def restore_database_backup(filename: str):
    try:
        backup_manager.restore_backup(filename)
        await broadcast_change("system", "restore")
        return {"success": True, "message": f"تمت استعادة النسخة الاحتياطية ({filename}) بنجاح"}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/backup/upload")
async def upload_database_backup(file: UploadFile = File(...)):
    if not file.filename.endswith(".db"):
        raise HTTPException(status_code=400, detail="يجب أن يكون الملف بصيغة قاعدة بيانات SQLite (.db)")
    backup_manager.ensure_backup_dir()
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    saved_filename = f"almanfath_uploaded_{timestamp}.db"
    saved_filepath = os.path.join(database.BACKUPS_DIR, saved_filename)
    with open(saved_filepath, "wb") as f:
        content = await file.read()
        f.write(content)
    # Restore from uploaded file
    backup_manager.restore_backup(saved_filename)
    await broadcast_change("system", "restore")
    return {"success": True, "message": "تم رفع واستعادة النسخة الاحتياطية بنجاح"}

@app.post("/api/reset-demo")
def reset_demo_data():
    conn = database.get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM customers")
    cursor.execute("DELETE FROM trips")
    database.seed_data(conn)
    conn.close()
    return {"success": True, "message": "تمت إعادة ضبط البيانات التجريبية بنجاح"}

@app.api_route("/", methods=["GET", "HEAD"], response_class=HTMLResponse)
def index():
    html_path = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(html_path):
        with open(html_path, "r", encoding="utf-8") as f:
            return f.read()
    return "<h1>نظام المنفذ المثالي - جاري تحميل الملفات...</h1>"

@app.api_route("/healthz", methods=["GET", "HEAD"])
def health_check():
    return {"status": "ok"}

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("server:app", host="0.0.0.0", port=port, reload=False)